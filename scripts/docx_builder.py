import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_header_footer(doc, title_text):
    for s in doc.sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)
        header = s.header
        hp = header.paragraphs[0]
        hp.text = f"AUDIO STREAMING ARCHITECTURE SPECIFICATION | {title_text}"
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hp.style.font.name = "Arial"
        hp.style.font.size = Pt(8.5)
        hp.style.font.color.rgb = RGBColor(120, 120, 120)

def create_document_base(title_text):
    doc = docx.Document()
    add_header_footer(doc, title_text)
    
    # Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Arial'
    normal_style.font.size = Pt(10)
    normal_style.font.color.rgb = RGBColor(33, 37, 41)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(4)
    
    return doc

def add_title(doc, title, subtitle="Hệ Thống Web Audio Phục Vụ Máy Chủ Tài Nguyên Hạn Chế"):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    run.font.name = "Arial"
    run.font.size = Pt(20)
    run.font.bold = True
    run.font.color.rgb = RGBColor(15, 42, 74) # Deep Navy
    
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run2 = p2.add_run(subtitle)
    run2.font.name = "Arial"
    run2.font.size = Pt(11)
    run2.font.italic = True
    run2.font.color.rgb = RGBColor(80, 90, 105)
    
    # Divider line
    p_div = doc.add_paragraph()
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_div = p_div.add_run("―" * 45)
    r_div.font.color.rgb = RGBColor(200, 205, 215)
    doc.add_paragraph()

def add_heading_1(doc, text):
    h = doc.add_heading(level=1)
    run = h.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = RGBColor(15, 42, 74)
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.space_after = Pt(4)
    return h

def add_heading_2(doc, text):
    h = doc.add_heading(level=2)
    run = h.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = RGBColor(31, 78, 121)
    h.paragraph_format.space_before = Pt(8)
    h.paragraph_format.space_after = Pt(3)
    return h

def add_heading_3(doc, text):
    h = doc.add_heading(level=3)
    run = h.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(10.5)
    run.font.bold = True
    run.font.color.rgb = RGBColor(60, 60, 60)
    h.paragraph_format.space_before = Pt(6)
    h.paragraph_format.space_after = Pt(2)
    return h

def add_callout(doc, text, tag="[FACT]"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
    
    bg_map = {
        "[FACT]": "EBF3FA",
        "[ANALYSIS]": "F0F4F8",
        "[RECOMMENDATION]": "EAF5EA",
        "[ESTIMATE]": "FEF8E7",
        "[WARNING]": "FDEDEC",
        "[LEGAL WARNING]": "FADBD8"
    }
    bg = bg_map.get(tag, "F4F6F7")
    set_cell_background(cell, bg)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run_tag = p.add_run(f"{tag} ")
    run_tag.bold = True
    if "WARNING" in tag:
        run_tag.font.color.rgb = RGBColor(180, 20, 20)
    elif tag == "[RECOMMENDATION]":
        run_tag.font.color.rgb = RGBColor(20, 110, 40)
    else:
        run_tag.font.color.rgb = RGBColor(15, 42, 74)
        
    run_text = p.add_run(text)
    run_text.font.size = Pt(9.5)
    doc.add_paragraph()

def add_code_block(doc, code_text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_background(cell, "212529")
    set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(code_text)
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(230, 235, 240)
    doc.add_paragraph()

def add_table_data(doc, headers, data, col_widths=None):
    table = doc.add_table(rows=len(data)+1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    # Header Row
    hdr_cells = table.rows[0].cells
    for i, h_text in enumerate(headers):
        hdr_cells[i].text = h_text
        set_cell_background(hdr_cells[i], "1F4E79")
        set_cell_margins(hdr_cells[i], top=80, bottom=80, left=100, right=100)
        p = hdr_cells[i].paragraphs[0]
        p.runs[0].font.bold = True
        p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
        p.runs[0].font.size = Pt(9)
        
    for r_idx, row_data in enumerate(data):
        row_cells = table.rows[r_idx+1].cells
        bg_col = "F9FAFB" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, cell_value in enumerate(row_data):
            row_cells[c_idx].text = str(cell_value)
            set_cell_background(row_cells[c_idx], bg_col)
            set_cell_margins(row_cells[c_idx], top=60, bottom=60, left=100, right=100)
            p = row_cells[c_idx].paragraphs[0]
            if len(p.runs) > 0:
                p.runs[0].font.size = Pt(8.5)
                
    if col_widths:
        for r in table.rows:
            for i, w in enumerate(col_widths):
                r.cells[i].width = Inches(w)
                
    doc.add_paragraph()
