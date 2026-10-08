import docx
import os
import glob

def extract_all():
    docs = sorted(glob.glob("docs_archive/*.docx"))
    out_dir = "docs_extracted"
    os.makedirs(out_dir, exist_ok=True)
    
    summary_lines = []
    
    for path in docs:
        basename = os.path.basename(path).replace(".docx", ".txt")
        out_path = os.path.join(out_dir, basename)
        doc = docx.Document(path)
        
        lines = []
        for p in doc.paragraphs:
            t = p.text.strip()
            if t:
                lines.append(t)
        for t_idx, table in enumerate(doc.tables):
            lines.append(f"\n--- TABLE {t_idx+1} ---")
            for row in table.rows:
                cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                lines.append(" | ".join(cells))
                
        with open(out_path, "w", encoding="utf-8") as f:
            f.write("\n".join(lines))
            
        summary_lines.append(f"{path} -> {len(lines)} lines")
        
    print("\n".join(summary_lines))

if __name__ == "__main__":
    extract_all()
